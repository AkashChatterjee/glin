"""Builds .xlsx fixtures programmatically for the Excel-input tests."""
import numpy as np


def make_workbook(path, sheets):
    """sheets: {name: (rows, merges)}; rows are lists written from A1."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (rows, merges) in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        for rng in merges:
            ws.merge_cells(rng)
    wb.save(path)
    return path


def churn_rows(n=200, seed=0):
    rng = np.random.RandomState(seed)
    ages = rng.randint(18, 80, size=n)
    plans = rng.choice(["basic", "premium"], size=n)
    return [
        [int(a), str(p), "Yes" if a < 40 else "No"] for a, p in zip(ages, plans)
    ]


def messy_churn_sheet(n=200):
    """Title banner, blank row, header, data, footnote."""
    rows = [
        ["Customer churn extract"],
        ["Source: internal"],
        [],
        ["age", "plan", "churn"],
        *churn_rows(n),
        [],
        ["Note: generated for tests"],
    ]
    return rows, []
